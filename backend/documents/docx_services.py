import re
from io import BytesIO

from django.core.exceptions import ValidationError
from docx import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from integrations.models import StorageProvider

from .models import Document
from .storage import DocumentStorageUnavailable, get_storage_backend

HEX_COLOR = re.compile(r"^#[0-9a-fA-F]{6}$")


def _rgb(value, fallback):
    selected = value if isinstance(value, str) and HEX_COLOR.fullmatch(value) else fallback
    return RGBColor.from_string(selected[1:])


def _add_logo(word, branding):
    logo_key = branding.get("logo_storage_key")
    provider_id = branding.get("logo_storage_provider_id")
    if not logo_key or not provider_id:
        return
    try:
        provider = StorageProvider.objects.get(pk=provider_id)
        stored = get_storage_backend(provider).open_stream(logo_key)
        image_data = stored.body.read()
        stored.body.close()
        paragraph = word.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.add_run().add_picture(BytesIO(image_data), width=Inches(2.15))
    except (StorageProvider.DoesNotExist, DocumentStorageUnavailable, OSError, ValueError):
        return


def render_generated_document_docx(document):
    if document.source_type != Document.SourceType.GENERATED or not document.rendered_content:
        raise ValidationError("Only generated documents can be exported as DOCX.")

    branding = document.rendered_branding or {}
    primary = _rgb(branding.get("primary"), "#8A5A00")
    accent = _rgb(branding.get("accent"), "#A86B00")
    word = WordDocument()
    section = word.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    _add_logo(word, branding)
    header = word.add_paragraph()
    header.paragraph_format.space_after = Pt(4)
    run = header.add_run(
        branding.get("header_text") or branding.get("brand_name") or "Evolve document"
    )
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = accent

    title = word.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run(document.title)
    run.bold = True
    run.font.size = Pt(21)
    run.font.color.rgb = primary

    meta = word.add_paragraph(
        f"{document.get_document_type_display()} · version {document.version_number}"
    )
    meta.paragraph_format.space_after = Pt(14)
    for run in meta.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    blocks = re.split(r"\n\s*---\s*\n", document.rendered_content)
    for block in blocks:
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if not lines:
            continue
        heading = word.add_paragraph()
        heading.paragraph_format.space_before = Pt(9)
        heading.paragraph_format.space_after = Pt(4)
        run = heading.add_run(lines[0])
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = primary
        if len(lines) > 1:
            body = word.add_paragraph("\n".join(lines[1:]))
            body.paragraph_format.space_after = Pt(7)
            for run in body.runs:
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0x20, 0x21, 0x24)

    if branding.get("footer_text"):
        footer = word.add_paragraph(branding["footer_text"])
        footer.paragraph_format.space_before = Pt(18)
        for run in footer.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    output = BytesIO()
    word.save(output)
    return output.getvalue()
